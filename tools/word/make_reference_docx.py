"""Build tools/word/reference.docx, the style template for the Word edition.

pandoc copies every style from this file into the book it writes, so all the
look of the Word edition lives here: fonts, heading colours, the smaller code
font, captions, page size and page numbers. Rerun after changing anything:

    python3 tools/word/make_reference_docx.py

Needs python-docx and the pandoc that ships with Quarto (for the starting file).
"""

import subprocess
from pathlib import Path

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "reference.docx"
TEAL = RGBColor(0x0F, 0x5F, 0x55)
INK = RGBColor(0x1F, 0x29, 0x33)
BODY, HEAD, CODE = "Georgia", "Calibri", "Consolas"


def start_file(quarto="quarto"):
    subprocess.run([quarto, "pandoc", "-o", str(OUT), "--print-default-data-file",
                    "reference.docx"], check=True)


def font(style, name, size, bold=None, italic=None, colour=None):
    f = style.font
    f.name = name
    f.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if colour is not None:
        f.color.rgb = colour


def shade(style, fill="F3F4F6"):
    ppr = style.element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(9)


def main(quarto="quarto"):
    start_file(quarto)
    d = docx.Document(str(OUT))
    st = d.styles

    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        if name in [s.name for s in st]:
            font(st[name], BODY, 11, colour=INK)
            pf = st[name].paragraph_format
            pf.line_spacing = 1.2
            pf.space_after = Pt(6)

    sizes = {1: 20, 2: 15, 3: 12.5, 4: 11.5}
    for level, size in sizes.items():
        s = st[f"Heading {level}"]
        font(s, HEAD, size, bold=True, colour=TEAL if level <= 2 else INK)
        s.paragraph_format.space_before = Pt(18 if level == 1 else 12)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
    # Page breaks before chapters are inserted explicitly by word.lua: many viewers
    # ignore a style-level "page break before".
    st["Heading 1"].paragraph_format.page_break_before = False

    for name in ("Title",):
        font(st[name], HEAD, 30, bold=True, colour=TEAL)
    font(st["Subtitle"], HEAD, 16, colour=INK)

    for name in ("Caption", "Image Caption", "Table Caption"):
        font(st[name], BODY, 9, italic=True, colour=INK)
        st[name].paragraph_format.space_after = Pt(10)

    # Code: smaller but readable, on a light grey band.
    names = [s.name for s in st]
    code = st["Source Code"] if "Source Code" in names else st.add_style(
        "Source Code", WD_STYLE_TYPE.PARAGRAPH)
    font(code, CODE, 8.5, colour=INK)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_after = Pt(0)
    shade(code)
    vchar = st["Verbatim Char"] if "Verbatim Char" in names else st.add_style(
        "Verbatim Char", WD_STYLE_TYPE.CHARACTER)
    font(vchar, CODE, 9, colour=INK)

    for section in d.sections:
        section.page_height, section.page_width = Cm(29.7), Cm(21.0)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(2.5))
        page_number_footer(section)

    # The starting file contains sample paragraphs; pandoc ignores the body, but
    # an empty body keeps the template clean.
    body = d.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    d.save(str(OUT))
    print(f"wrote {OUT}")


if __name__ == "__main__":
    import sys
    main(sys.argv[1] if len(sys.argv) > 1 else "quarto")
